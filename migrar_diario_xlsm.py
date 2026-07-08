# =============================================================================
# migrar_diario_xlsm.py — Import único e unidirecional do diário Excel
# =============================================================================
# Lê `Diário de Trade 2026 MACRO.xlsm` (abas Diario + Sim) UMA VEZ, extrai as
# imagens embutidas via rich-value (picture-in-cell), normaliza os campos
# livres (Janela/Estratégia/Sala) e grava em `diario_trades.parquet` via
# `persistencia.py`/`diario_trade.py`. Nunca escreve de volta no xlsm, nunca
# precisa rodar de novo depois de aprovado (ver `--forcar`).
#
# Uso:
#   python migrar_diario_xlsm.py --dry-run   # SEMPRE rodar assim primeiro —
#                                             # não grava nada, só imprime os
#                                             # value_counts dos mapeamentos
#                                             # pra revisão do Jesué.
#   python migrar_diario_xlsm.py              # escrita real (recusa se já
#                                              # houver linhas importadas
#                                              # antes, a menos que --forcar)
#   python migrar_diario_xlsm.py --forcar      # permite reimportar por cima
#
# Mesmo espírito dos scripts de investigação já existentes no projeto
# (`investiga_*.py`) — raro, one-shot, não faz parte do pipeline em produção.
# =============================================================================

from __future__ import annotations

import argparse
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import date, datetime

import openpyxl
import pandas as pd

import config_backtest as cfg
import diario_trade as dt
import persistencia

NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
NS_XLRD = "{http://schemas.microsoft.com/office/spreadsheetml/2017/richdata}"
NS_RV = "{http://schemas.microsoft.com/office/spreadsheetml/2017/richdata}"
NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
NS_RVREL = "{http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel}"
NS_PKGREL = "{http://schemas.openxmlformats.org/package/2006/relationships}"

# Baseline conhecido nesta migração (ver plano) — só informativo: o xlsm pode
# ter crescido desde então (novos trades lançados antes de rodar isto). Não é
# usado como validação rígida, só como referência impressa pro Jesué.
BASELINE_IMAGENS_DIARIO = 117
BASELINE_IMAGENS_SIM = 12

# -----------------------------------------------------------------------------
# MAPEAMENTOS DE NORMALIZAÇÃO — revisar com Jesué via --dry-run antes de gravar
# Chave = valor cru, em minúsculas e sem espaço extra nas pontas.
# -----------------------------------------------------------------------------

MAPA_JANELA = {
    "abertura (9:00 - 9:30)": "abertura",
    "prime (9:30 - 11:00)": "prime",
    "almoço (11:00 - 13:00)": "almoco",
    "almoco (11:00 - 13:00)": "almoco",
    "almoço (12:00 - 13:00)": "almoco",
    "tarde (13:00 - 15:00)": "tarde",
    "final (15:00 - 17:00)": "final",
}

MAPA_SALA_ANALISE = {
    "btv igor": "btv_igor",
    "b.trader igor": "btv_igor",
    "btv marquinhos": "btv_marquinhos",
    "b.trader marquinhos": "btv_marquinhos",
    "btv kita": "btv_kita",
    "analise propria": "analise_propria",
    "análise própria": "analise_propria",
}

# Valor = lista de tags canônicas (`diario_trade.ESTRATEGIA_TAGS_CANONICAS`)
# — strings compostas (ex. "LTB/Canal") viram mais de uma tag.
MAPA_ESTRATEGIA = {
    "suporte\\resitência": ["suporte_resistencia"],
    "suporte/resistencia": ["suporte_resistencia"],
    "bounce snr": ["suporte_resistencia"],
    "inversao de polaridade": ["inversao_polaridade"],
    "inversão de polaridade": ["inversao_polaridade"],
    "mudança de polaridade": ["inversao_polaridade"],
    "pullback mme21": ["pullback_mme21"],
    "entrada mme21": ["pullback_mme21"],
    "mme21": ["pullback_mme21"],
    "médias": ["medias_moveis"],
    "bandeira": ["bandeira"],
    "canal": ["canal"],
    "rompimento pivot": ["rompimento"],
    "pullback rompimento": ["rompimento", "reteste"],
    "trowback": ["reteste"],
    "lta\\ltb": ["lta_ltb"],
    "ltb/canal": ["lta_ltb", "canal"],
    "ltb\\canal": ["lta_ltb", "canal"],
    "retração fibo": ["fibonacci"],
    "oco": ["ombro_cabeca_ombro"],
    # Sem correspondência clara — vira ["outro"], texto original preservado
    # em `estrategia_raw`. REVISAR no --dry-run antes de aprovar.
    "espadada": ["outro"],
    "figura gráfica": ["outro"],
    "kita": ["outro"],           # provável erro de digitação (nome de mentor
                                  # no campo de estratégia) — conferir a linha
    "pullback vwap": ["outro"],
    "fechamento de gap": ["outro"],
    "jacaré": ["outro"],
    "pullback": ["outro"],
}


# -----------------------------------------------------------------------------
# EXTRAÇÃO DE IMAGENS RICH-VALUE (picture-in-cell)
# -----------------------------------------------------------------------------
# Cadeia completa (testada célula a célula contra o arquivo real antes de
# escrever este script — ver plano): célula `vm="M"` (1-based) -> M-ésimo
# <bk> de <valueMetadata> em xl/metadata.xml -> <rc t="T" v="V"/> (T = índice
# 1-based do metadataType, V = índice 0-based dentro do <futureMetadata>
# daquele tipo) -> V-ésimo <bk> de <futureMetadata name="XLRICHVALUE"> ->
# <xlrd:rvb i="K"/> -> K-ésimo <rv> de xl/richData/rdrichvalue.xml -> 1º <v>
# = índice L em xl/richData/richValueRel.xml -> L-ésimo <rel r:id="rIdX"/>
# (0-based) -> xl/richData/_rels/richValueRel.xml.rels resolve rIdX ->
# xl/media/imageY.png.

def _extrair_imagens_richvalue(caminho_xlsm) -> dict[int, bytes]:
    with zipfile.ZipFile(caminho_xlsm) as z:
        metadata_root = ET.fromstring(z.read("xl/metadata.xml"))

        value_metadata = metadata_root.find(f"{NS_MAIN}valueMetadata")
        vm_records = []
        for bk in value_metadata.findall(f"{NS_MAIN}bk"):
            rc = bk.find(f"{NS_MAIN}rc")
            vm_records.append((int(rc.get("t")), int(rc.get("v"))))

        metadata_types = metadata_root.find(f"{NS_MAIN}metadataTypes")
        nomes_tipos = [mt.get("name") for mt in metadata_types.findall(f"{NS_MAIN}metadataType")]
        indice_richvalue = nomes_tipos.index("XLRICHVALUE") + 1  # 1-based

        fm_richvalue = next(
            fm for fm in metadata_root.findall(f"{NS_MAIN}futureMetadata")
            if fm.get("name") == "XLRICHVALUE"
        )
        bk_para_rv_idx = [
            int(bk.find(f".//{NS_XLRD}rvb").get("i"))
            for bk in fm_richvalue.findall(f"{NS_MAIN}bk")
        ]

        rv_root = ET.fromstring(z.read("xl/richData/rdrichvalue.xml"))
        rv_para_relrel_idx = [
            int(rv.find(f"{NS_RV}v").text) for rv in rv_root.findall(f"{NS_RV}rv")
        ]

        relrel_root = ET.fromstring(z.read("xl/richData/richValueRel.xml"))
        relrel_para_rid = [
            rel.get(f"{NS_R}id") for rel in relrel_root.findall(f"{NS_RVREL}rel")
        ]

        rels_root = ET.fromstring(z.read("xl/richData/_rels/richValueRel.xml.rels"))
        rid_para_target = {
            rel.get("Id"): rel.get("Target")
            for rel in rels_root.findall(f"{NS_PKGREL}Relationship")
        }

        imagens: dict[int, bytes] = {}
        for vm_1based, (tipo_idx, v_idx) in enumerate(vm_records, start=1):
            if tipo_idx != indice_richvalue:
                continue  # metadado de outro tipo (ex. array dinâmico), não é imagem
            rv_idx = bk_para_rv_idx[v_idx]
            relrel_idx = rv_para_relrel_idx[rv_idx]
            rid = relrel_para_rid[relrel_idx]
            caminho_interno = "xl/" + rid_para_target[rid].replace("../", "")
            imagens[vm_1based] = z.read(caminho_interno)

    return imagens


def _mapa_sheet_para_xml(caminho_xlsm) -> dict[str, str]:
    """Resolve nome de aba -> caminho do XML da planilha dentro do zip, sem hardcode de sheetN."""
    with zipfile.ZipFile(caminho_xlsm) as z:
        wb_root = ET.fromstring(z.read("xl/workbook.xml"))
        rels_root = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        rid_para_target = {
            rel.get("Id"): rel.get("Target")
            for rel in rels_root.findall(f"{NS_PKGREL}Relationship")
        }
        sheets = wb_root.find(f"{NS_MAIN}sheets")
        return {
            sheet.get("name"): "xl/" + rid_para_target[sheet.get(f"{NS_R}id")]
            for sheet in sheets.findall(f"{NS_MAIN}sheet")
        }


def _extrair_vm_por_linha(caminho_xlsm, caminho_sheet_xml: str, coluna: str) -> dict[int, int]:
    """Regex sobre a coluna de Print da aba: linha -> valor de `vm` (1-based)."""
    with zipfile.ZipFile(caminho_xlsm) as z:
        conteudo = z.read(caminho_sheet_xml).decode("utf-8", errors="replace")
    padrao = re.compile(rf'<c r="{coluna}(\d+)"[^>]*vm="(\d+)"')
    return {int(linha): int(vm) for linha, vm in padrao.findall(conteudo)}


# -----------------------------------------------------------------------------
# NORMALIZAÇÃO DE CAMPOS
# -----------------------------------------------------------------------------

def _norm(texto) -> str:
    return str(texto).strip().lower() if texto is not None else ""


def _parse_data(valor, nao_mapeados: Counter | None = None, campo: str = "data"):
    """
    Retorna None (em vez de lançar) pra valores que não são data de verdade —
    células como '-' aparecem legitimamente (ex. `Data Replay` de um trade
    'Ao Vivo', que não teve replay) — mas registra em `nao_mapeados` pra
    aparecer no resumo do --dry-run, nunca falha em silêncio.
    """
    if valor is None:
        return None
    if isinstance(valor, (datetime, date)):
        return pd.Timestamp(valor).date()
    if isinstance(valor, str):
        texto = valor.strip()
        if texto in ("", "-"):
            return None
        for sep in ("\\", "/"):
            partes = texto.split(sep)
            if len(partes) == 3:
                try:
                    d, m, a = (int(p) for p in partes)
                    return date(a, m, d)
                except ValueError:
                    continue
    if nao_mapeados is not None:
        nao_mapeados[f"{campo} não reconhecida: {valor!r}"] += 1
    return None


def _mapear_janela(valor, nao_mapeados: Counter) -> str | None:
    if valor is None:
        return None
    chave = _norm(valor)
    if chave in MAPA_JANELA:
        return MAPA_JANELA[chave]
    nao_mapeados[f"janela: {valor!r}"] += 1
    return None


def _mapear_sala(valor, nao_mapeados: Counter) -> str | None:
    if valor is None:
        return None
    chave = _norm(valor)
    if chave in MAPA_SALA_ANALISE:
        return MAPA_SALA_ANALISE[chave]
    nao_mapeados[f"sala: {valor!r}"] += 1
    return "outro"


def _mapear_estrategia(valor, nao_mapeados: Counter) -> list[str]:
    if valor is None:
        return []
    chave = _norm(valor)
    if chave in MAPA_ESTRATEGIA:
        return MAPA_ESTRATEGIA[chave]
    nao_mapeados[f"estrategia: {valor!r}"] += 1
    return ["outro"]


# -----------------------------------------------------------------------------
# LEITURA DAS ABAS
# -----------------------------------------------------------------------------

def _ultima_linha_populada(ws, col_a: int, col_c: int, linha_inicial: int, tolerancia_vazias: int = 5) -> int:
    ultima = linha_inicial - 1
    vazias_seguidas = 0
    linha = linha_inicial
    while True:
        a = ws.cell(row=linha, column=col_a).value
        c = ws.cell(row=linha, column=col_c).value
        if a is None and c is None:
            vazias_seguidas += 1
            if vazias_seguidas > tolerancia_vazias:
                break
        else:
            vazias_seguidas = 0
            ultima = linha
        linha += 1
    return ultima


def _ler_linhas_diario(ws, imagens_por_vm: dict[int, int], nao_mapeados: Counter) -> list[dict]:
    ultima = _ultima_linha_populada(ws, col_a=1, col_c=3, linha_inicial=11)
    linhas = []
    for linha in range(11, ultima + 1):
        v = lambda col: ws.cell(row=linha, column=col).value
        data_op = v(1)
        ativo = v(3)
        if data_op is None and ativo is None:
            continue  # linha em branco no meio do range (não deveria acontecer, mas não falha por isso)

        tipo_saida_raw = v(12)
        is_simulado = _norm(tipo_saida_raw) == "simulador"

        linhas.append({
            "data_op": _parse_data(data_op, nao_mapeados, "data_op (Diario)"),
            "ativo": ativo,
            "janela_pregao": _mapear_janela(v(2), nao_mapeados),
            "tempo_grafico": v(4) or "5 min.",
            "estrategia_raw": v(5),
            "estrategia_tags": _mapear_estrategia(v(5), nao_mapeados),
            "sala_analise_raw": v(6),
            "sala_analise": _mapear_sala(v(6), nao_mapeados),
            "alvo_pontos": v(7),
            "stop_pontos": v(8),
            "lote": v(9) or 1,
            "tipo_operacao": _norm(v(10)) or None,
            "preco_entrada": v(11),
            "tipo_saida": None if is_simulado else _norm(tipo_saida_raw) or None,
            "preco_saida": v(13),
            "resultado_pontos": v(14),
            "saldo_operacao_reais": v(15),
            "resumo_operacional": v(16),
            "obs": v(18),
            "is_simulated": is_simulado,
            "conta": "simulador" if is_simulado else "real",
            "data_replay": None,
            "_vm_imagem": imagens_por_vm.get(linha),
            "_origem_sheet": "Diario",
            "_origem_linha": linha,
        })
    return linhas


def _ler_linhas_sim(ws, imagens_por_vm: dict[int, int], nao_mapeados: Counter) -> list[dict]:
    ultima = _ultima_linha_populada(ws, col_a=1, col_c=4, linha_inicial=2)
    linhas = []
    for linha in range(2, ultima + 1):
        v = lambda col: ws.cell(row=linha, column=col).value
        data_op = v(1)
        ativo = v(4)
        if data_op is None and ativo is None:
            continue

        mercado = v(7)  # 'Replay'/'Ao Vivo' — sem coluna própria no schema
        obs_original = v(18)
        obs = f"[Mercado: {mercado}] {obs_original}" if mercado else obs_original

        linhas.append({
            "data_op": _parse_data(data_op, nao_mapeados, "data_op (Sim)"),
            "ativo": ativo,
            "janela_pregao": _mapear_janela(v(3), nao_mapeados),
            "tempo_grafico": v(5) or "5 min.",
            "estrategia_raw": v(8),
            "estrategia_tags": _mapear_estrategia(v(8), nao_mapeados),
            "sala_analise_raw": v(9),
            "sala_analise": _mapear_sala(v(9), nao_mapeados),
            "alvo_pontos": v(10),
            "stop_pontos": v(11),
            "lote": v(12) or 1,
            "tipo_operacao": _norm(v(13)) or None,
            "preco_entrada": None,
            "tipo_saida": None,
            "preco_saida": None,
            "resultado_pontos": v(14),
            "saldo_operacao_reais": v(15),
            "resumo_operacional": v(17),
            "obs": obs,
            "is_simulated": True,
            "conta": "simulador",
            "data_replay": _parse_data(v(2), nao_mapeados, "data_replay (Sim)"),
            "_vm_imagem": imagens_por_vm.get(linha),
            "_origem_sheet": "Sim",
            "_origem_linha": linha,
        })
    return linhas


# -----------------------------------------------------------------------------
# ORQUESTRAÇÃO
# -----------------------------------------------------------------------------

def _montar_dataframe(caminho_xlsm) -> tuple[pd.DataFrame, dict[int, bytes], Counter]:
    nao_mapeados: Counter = Counter()
    mapa_sheets = _mapa_sheet_para_xml(caminho_xlsm)

    imagens = _extrair_imagens_richvalue(caminho_xlsm)
    vm_diario = _extrair_vm_por_linha(caminho_xlsm, mapa_sheets["Diario"], "Q")
    vm_sim = _extrair_vm_por_linha(caminho_xlsm, mapa_sheets["Sim"], "S")

    wb = openpyxl.load_workbook(caminho_xlsm, read_only=True, data_only=True, keep_vba=False)
    linhas_diario = _ler_linhas_diario(wb["Diario"], vm_diario, nao_mapeados)
    linhas_sim = _ler_linhas_sim(wb["Sim"], vm_sim, nao_mapeados)
    todas = linhas_diario + linhas_sim

    # Ordena cronologicamente (Diario antes de Sim em empate de data, só por
    # estabilidade determinística) e atribui id_trade sequencial.
    todas.sort(key=lambda r: (r["data_op"] or date.min, r["_origem_sheet"], r["_origem_linha"]))
    for i, linha in enumerate(todas, start=1):
        linha["id_trade"] = i

    df = pd.DataFrame(todas)
    df["fonte_registro"] = dt.FONTE_IMPORTADO
    df["criado_em"] = datetime.now()
    df["contrato"] = None
    df["datetime_entrada"] = None
    df["datetime_saida"] = None
    df["nivel_sr_referencia"] = None
    df["nivel_sr_tipo"] = None
    df["nivel_sr_forte"] = None
    df["tags_setup_erro"] = [[] for _ in range(len(df))]
    df["imagem_arquivo"] = None

    return df, imagens, nao_mapeados


def _imprimir_resumo_dry_run(df: pd.DataFrame, imagens: dict[int, bytes], nao_mapeados: Counter):
    print(f"\n=== RESUMO --dry-run ({len(df)} trades encontrados) ===\n")
    print(f"Intervalo de datas: {df['data_op'].min()} a {df['data_op'].max()}")
    print(f"Imagens resolvidas via rich-value: {len(imagens)} "
          f"(baseline conhecido: {BASELINE_IMAGENS_DIARIO + BASELINE_IMAGENS_SIM} — "
          f"diferença é esperada se o xlsm cresceu desde o mapeamento original)")
    print(f"Trades com imagem: {df['_vm_imagem'].notna().sum()} / {len(df)}")
    print(f"Trades simulados (is_simulated=True): {df['is_simulated'].sum()}")

    for coluna in ["ativo", "janela_pregao", "sala_analise", "tipo_operacao", "tipo_saida"]:
        print(f"\n--- value_counts: {coluna} ---")
        print(df[coluna].value_counts(dropna=False).to_string())

    print("\n--- value_counts: estrategia_tags (achatado) ---")
    todas_tags = [tag for tags in df["estrategia_tags"] for tag in tags]
    print(pd.Series(todas_tags).value_counts().to_string())

    if nao_mapeados:
        print(f"\n!!! {sum(nao_mapeados.values())} valores SEM mapeamento explícito "
              f"(caíram em 'outro' ou None) — REVISAR MAPA_* antes de aprovar:")
        for chave, n in nao_mapeados.most_common():
            print(f"  {n:3d}  {chave}")
    else:
        print("\nTodos os valores categóricos tiveram mapeamento explícito.")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dry-run", action="store_true",
                         help="Só imprime o resumo/mapeamento, não grava nada em disco.")
    parser.add_argument("--forcar", action="store_true",
                         help="Permite rodar de novo mesmo se já houver linhas importadas.")
    args = parser.parse_args()

    caminho_xlsm = cfg.CAMINHO_DIARIO_XLSM
    if not caminho_xlsm.exists():
        print(f"ERRO: arquivo não encontrado: {caminho_xlsm}")
        sys.exit(1)

    if not args.dry_run and not args.forcar:
        existente = persistencia.carregar_diario()
        if existente is not None and (existente["fonte_registro"] == dt.FONTE_IMPORTADO).any():
            print("ERRO: já existem linhas importadas em diario_trades.parquet. "
                  "Rode com --forcar se quiser reimportar por cima (vai duplicar "
                  "a menos que o parquet seja apagado antes).")
            sys.exit(1)

    print(f"Lendo {caminho_xlsm} ...")
    df, imagens, nao_mapeados = _montar_dataframe(caminho_xlsm)

    if args.dry_run:
        _imprimir_resumo_dry_run(df, imagens, nao_mapeados)
        print("\n--dry-run: nada foi gravado em disco.")
        return

    with persistencia.lock_diario() as adquirida:
        if not adquirida:
            print("ERRO: diário está travado por outro processo — tente de novo em instantes.")
            sys.exit(1)

        # Salva as imagens ANTES de descartar a coluna auxiliar `_vm_imagem`.
        for _, linha in df.iterrows():
            vm = linha["_vm_imagem"]
            if pd.isna(vm):
                continue
            conteudo = imagens.get(int(vm))
            if conteudo is None:
                print(f"AVISO: trade id={linha['id_trade']} referencia vm={vm} sem imagem resolvida — pulando.")
                continue
            nome_arquivo = dt.salvar_imagem(conteudo, int(linha["id_trade"]), linha["ativo"], linha["data_op"])
            df.loc[df["id_trade"] == linha["id_trade"], "imagem_arquivo"] = nome_arquivo

        df_final = df.drop(columns=["_vm_imagem", "_origem_sheet", "_origem_linha"])
        df_final = df_final[dt.COLUNAS_DIARIO]
        persistencia.salvar_diario(df_final)

    print(f"\nMigração concluída: {len(df_final)} trades gravados em diario_trades.parquet.")
    print(f"Imagens salvas em {cfg.PASTA_DIARIO_IMAGENS}")


if __name__ == "__main__":
    main()
